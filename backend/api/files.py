from fastapi import APIRouter, HTTPException, UploadFile, File
import io

from backend.models.schemas import FileGenerateRequest, FileGenerateResponse
from backend.services.file_service import generate_file

router = APIRouter()


@router.post("/files", response_model=FileGenerateResponse)
async def files_endpoint(req: FileGenerateRequest):
    """Gera PDF, DOCX, XLSX ou PPTX e faz upload ao Supabase."""
    try:
        url, message = await generate_file(req.type, req.title, req.content)
        return FileGenerateResponse(url=url, message=message)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/extract-text")
async def extract_text_endpoint(file: UploadFile = File(...)):
    """Extrai texto de PDFs, DOCX ou XLSX para uso no chat."""
    filename = file.filename.lower()
    
    try:
        content = await file.read()
        
        if filename.endswith(".pdf"):
            import PyPDF2
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() for page in pdf_reader.pages if page.extract_text())
            return {"text": text}
            
        elif filename.endswith(".docx"):
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            return {"text": text}
            
        elif filename.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            text_lines = []
            for sheet in wb.worksheets:
                text_lines.append(f"--- Planilha: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    # Filter out purely None rows
                    row_vals = [str(cell) if cell is not None else "" for cell in row]
                    if any(row_vals):
                        text_lines.append("\t".join(row_vals))
            return {"text": "\n".join(text_lines)}
            
        else:
            # Fallback for txt, csv, md
            return {"text": content.decode("utf-8", errors="replace")}
            
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Erro ao extrair texto: {str(e)}")
